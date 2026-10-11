from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
import os

# Create presentation
prs = Presentation()

# Layouts (0: Title, 1: Title and Content, 3: Two Content, 5: Title Only)
title_slide_layout = prs.slide_layouts[0]
content_slide_layout = prs.slide_layouts[1]
two_content_layout = prs.slide_layouts[3]

# ---------------------------------------------------
# SLIDE 1: Title Slide
# ---------------------------------------------------
slide1 = prs.slides.add_slide(title_slide_layout)
title = slide1.shapes.title
subtitle = slide1.placeholders[1]

title.text = "Weather-Driven Business Impact: A Multi-Source EDA & Econometric Analysis"
title.text_frame.paragraphs[0].font.size = Pt(36)
title.text_frame.paragraphs[0].font.bold = True

subtitle.text = "Course Project Review 1\nExploratory Data Analysis\n\nTeam Members: [Insert Names]"
subtitle.text_frame.paragraphs[0].font.size = Pt(24)

# ---------------------------------------------------
# SLIDE 2: Abstract (The Problem & Solution)
# ---------------------------------------------------
slide2 = prs.slides.add_slide(content_slide_layout)
title2 = slide2.shapes.title
title2.text = "Abstract: Problem & Solution"

body_shape2 = slide2.placeholders[1]
tf2 = body_shape2.text_frame
tf2.text = "The Problem:"
p = tf2.add_paragraph()
p.text = "Standard business analysis often confuses normal calendar trends (like busy weekends) with actual weather impacts."
p.level = 1

p = tf2.add_paragraph()
p.text = "The Solution:"
p.level = 0
p = tf2.add_paragraph()
p.text = "We are building an interactive Exploratory Data Analysis (EDA) platform to isolate the true financial impact of weather across 16+ commercial sectors."
p.level = 1

p = tf2.add_paragraph()
p.text = "The Goal:"
p.level = 0
p = tf2.add_paragraph()
p.text = "To visually map the 'Domino Effect' of weather (e.g., proving how heavy rain suppresses bike-sharing, surges rideshare apps, lowers dine-in restaurant traffic, and spikes food delivery)."
p.level = 1

# Try adding an image to Slide 2 if available
img_path_1 = r"C:\Users\vyasv\.gemini\antigravity\brain\b3a50ffe-df6a-4036-b560-a035c17603f4\.user_uploaded\media_1787156488189.png"
if os.path.exists(img_path_1):
    try:
        # Resize and place image on the right
        slide2.shapes.add_picture(img_path_1, Inches(5.5), Inches(2.5), width=Inches(4.0))
        # Shrink text box width to fit image
        body_shape2.width = Inches(5.0)
    except:
        pass


# ---------------------------------------------------
# SLIDE 3: Dataset Details (Sources & Metrics)
# ---------------------------------------------------
slide3 = prs.slides.add_slide(two_content_layout)
title3 = slide3.shapes.title
title3.text = "Dataset Details: Sources & Metrics"

left_body = slide3.placeholders[1]
tf3_left = left_body.text_frame
tf3_left.text = "Weather & Climate Data (The Cause):"
p = tf3_left.add_paragraph()
p.text = "Sources: NOAA and OpenWeather APIs."
p.level = 1
p = tf3_left.add_paragraph()
p.text = "Metrics: Hourly dry bulb temperature, precipitation (mm/hr), wind speed, and snowfall."
p.level = 1

right_body = slide3.placeholders[2]
tf3_right = right_body.text_frame
tf3_right.text = "Business & Economic Data (The Effect):"
p = tf3_right.add_paragraph()
p.text = "Sources: Stochastic Data Synthesis (creating highly realistic mock datasets based on real-world logic since corporate data is confidential)."
p.level = 1
p = tf3_right.add_paragraph()
p.text = "Coverage: 16+ sectors including Logistics (Uber/DoorDash), Hospitality, Retail, and Infrastructure (Power Grid)."
p.level = 1

# Try adding a different image at bottom of slide 3
img_path_2 = r"C:\Users\vyasv\.gemini\antigravity\brain\b3a50ffe-df6a-4036-b560-a035c17603f4\.user_uploaded\media_1787505258302.png"
if os.path.exists(img_path_2):
    try:
        slide3.shapes.add_picture(img_path_2, Inches(1.0), Inches(5.0), height=Inches(2.2))
    except:
        pass


# ---------------------------------------------------
# SLIDE 4: Dataset Details (EDA Challenges & Wrangling)
# ---------------------------------------------------
slide4 = prs.slides.add_slide(content_slide_layout)
title4 = slide4.shapes.title
title4.text = "EDA Challenges & Data Wrangling"

body_shape4 = slide4.placeholders[1]
tf4 = body_shape4.text_frame
tf4.text = "The Clock Problem (Timezones):"
p = tf4.add_paragraph()
p.text = "Weather sensors record in global UTC time, while businesses record in local time. We programmatically aligned these, accounting for Daylight Saving Time."
p.level = 1

p = tf4.add_paragraph()
p.text = "The Stopwatch Problem (Granularity):"
p.level = 0
p = tf4.add_paragraph()
p.text = "Business sales log every 15 minutes, but weather is hourly. We utilized Pandas resampling to stack 15-minute chunks into 1-hour blocks for side-by-side analysis."
p.level = 1

p = tf4.add_paragraph()
p.text = "The Broken Thermometer (Missing Data):"
p.level = 0
p = tf4.add_paragraph()
p.text = "When weather sensors drop offline, we used cubic spline interpolation (mathematical dot-connecting) to safely estimate missing temperatures without breaking our time-series."
p.level = 1


# ---------------------------------------------------
# SLIDE 5: Implementation RoadMap
# ---------------------------------------------------
slide5 = prs.slides.add_slide(content_slide_layout)
title5 = slide5.shapes.title
title5.text = "Implementation RoadMap"

body_shape5 = slide5.placeholders[1]
tf5 = body_shape5.text_frame

p = tf5.paragraphs[0]
p.text = "Phase 1: Data Sourcing (Completed) - Setup API integrations and synthesize proxy business datasets."
p = tf5.add_paragraph()
p.text = "Phase 2: Data Wrangling (Completed) - Cleanse data, align timezones, handle missing values, and merge into Parquet files."
p = tf5.add_paragraph()
p.text = "Phase 3: Exploratory Data Analysis (In Progress) - Visualize non-linear tipping points and cross-sector ripple effects."
p = tf5.add_paragraph()
p.text = "Phase 4: Econometric Modeling (Upcoming) - Apply Two-Way Fixed Effects (TWFE) to separate calendar trends from weather anomalies."
p = tf5.add_paragraph()
p.text = "Phase 5: Dashboard Deployment (Upcoming) - Finalize the interactive Streamlit dashboard with a live scenario simulator."

# Save PPTX
ppt_path = r"C:\Users\vyasv\.gemini\antigravity\scratch\weather-business-impact-eda\Course_Project_Review_1.pptx"
prs.save(ppt_path)
print(f"Presentation saved successfully to: {ppt_path}")
